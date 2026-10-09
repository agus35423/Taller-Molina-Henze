from django.contrib import messages
from django.contrib.auth import authenticate, login, logout
from django.contrib.auth.decorators import login_required
from django.contrib.auth.forms import AuthenticationForm
from django.contrib.auth.models import User
from django.core.mail import send_mail
from django.http import HttpResponse
from django.shortcuts import get_object_or_404, redirect, render
from django.utils import timezone


from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

from .forms import RegistroForm, SolicitarTurnoForm
from django.urls import reverse
from .models import Turno, Trabajo, OfertaTurno

# ==========================================================
# INICIO
# ==========================================================

def inicio(request):

    trabajos = Trabajo.objects.all()[:6]

    return render(
        request,
        "tallermami/inicio.html",
        {
            "trabajos": trabajos,
        }
    )


# ==========================================================
# LOGIN
# ==========================================================

def login_view(request):

    if request.user.is_authenticated:
        return redirect("inicio")

    if request.method == "POST":

        form = AuthenticationForm(
            request,
            data=request.POST
        )

        if form.is_valid():

            usuario = form.get_user()

            login(
                request,
                usuario
            )

            messages.success(
                request,
                "Iniciaste sesión correctamente."
            )

            return redirect("inicio")

    else:

        form = AuthenticationForm()

    return render(
        request,
        "tallermami/login.html",
        {
            "form": form
        }
    )


# ==========================================================
# REGISTRO
# ==========================================================

def registro(request):

    if request.user.is_authenticated:
        return redirect("inicio")

    if request.method == "POST":

        form = RegistroForm(request.POST)

        if form.is_valid():

            usuario = form.save()

            login(
                request,
                usuario
            )

            messages.success(
                request,
                "Tu cuenta fue creada correctamente."
            )

            return redirect("inicio")

    else:

        form = RegistroForm()

    return render(
        request,
        "tallermami/registro.html",
        {
            "form": form
        }
    )


# ==========================================================
# LOGOUT
# ==========================================================

def logout_view(request):

    logout(request)

    messages.success(
        request,
        "Sesión cerrada correctamente."
    )

    return redirect("inicio")


# ==========================================================
# SOLICITAR TURNO
# ==========================================================

def solicitar_turno(request):

    if request.method == "POST":

        form = SolicitarTurnoForm(
            request.POST
        )

        if form.is_valid():

            turno = form.save(
                commit=False
            )

            # Usuario autenticado
            if request.user.is_authenticated:

                turno.usuario = request.user

            # Si no está autenticado, usamos un usuario
            # técnico para poder guardar el turno.
            else:

                usuario, created = User.objects.get_or_create(
                    username="clientes"
                )

                if created:

                    usuario.set_unusable_password()

                    usuario.save()

                turno.usuario = usuario


            turno.save()


            # ==================================================
            # GUARDAR EN EXCEL
            # ==================================================

            try:

                guardar_turno_en_excel(turno)

            except Exception as error:

                print(
                    f"Error al guardar el turno en Excel: {error}"
                )


            # ==================================================
            # ENVIAR EMAIL
            # ==================================================

            email_cliente = form.cleaned_data.get(
                "email"
            )

            nombre_cliente = form.cleaned_data.get(
                "nombre"
            )

            if email_cliente:

                try:

                    send_mail(

                        subject="Solicitud de turno - Taller Molina Henze",

                        message=(
                            f"Hola {nombre_cliente},\n\n"
                            "Recibimos correctamente tu solicitud de turno.\n\n"
                            f"Vehículo: {turno.vehiculo}\n"
                            f"Patente: {turno.patente}\n"
                            f"Servicio: {turno.servicio}\n"
                            f"Fecha: {turno.fecha_turno.strftime('%d/%m/%Y %H:%M')}\n\n"
                            "Te esperamos en Taller Molina Henze.\n\n"
                            "Saludos."
                        ),

                        from_email=None,

                        recipient_list=[
                            email_cliente
                        ],

                        fail_silently=True
                    )

                except Exception as error:

                    print(
                        f"Error al enviar email: {error}"
                    )


            return redirect(
                "turno_exitoso",
                turno_id=turno.id
            )

    else:

        form = SolicitarTurnoForm()


    return render(
        request,
        "tallermami/solicitar_turno.html",
        {
            "form": form
        }
    )


# ==========================================================
# TURNO EXITOSO
# ==========================================================

def turno_exitoso(request, turno_id):

    turno = get_object_or_404(
        Turno,
        id=turno_id
    )

    return render(
        request,
        "tallermami/turno_exitoso.html",
        {
            "turno": turno
        }
    )


# ==========================================================
# MIS TURNOS
# ==========================================================




@login_required
def mis_turnos(request):

    turnos = Turno.objects.filter(
        usuario=request.user
    ).exclude(
        estado="cancelado"
    ).select_related(
        "servicio"
    )

    return render(
        request,
        "tallermami/mis_turnos.html",
        {
            "turnos": turnos
        }
    )



# ==========================================================
# CANCELAR TURNO
# ==========================================================

@login_required
def cancelar_turno(request, turno_id):

    turno = get_object_or_404(
        Turno,
        id=turno_id,
        usuario=request.user
    )

    # Solo permitimos cancelar turnos que todavía estén activos
    if turno.estado == "cancelado":
        messages.error(
            request,
            "Este turno ya está cancelado."
        )
        return redirect("mis_turnos")

    # Guardamos la fecha y hora que quedó libre
    fecha_liberada = turno.fecha_turno

    # Cancelamos el turno
    turno.estado = "cancelado"
    turno.save(update_fields=["estado"])

    # ==========================================================
    # BUSCAR CLIENTE PARA OFRECERLE EL HORARIO
    # ==========================================================

    turno_para_ofrecer = (
        Turno.objects
        .filter(
            fecha_turno__gt=timezone.now(),
            estado="pendiente"
        )
        .exclude(
            id=turno.id
        )
        .select_related(
            "usuario",
            "servicio"
        )
        .order_by("fecha_turno")
        .first()
    )

    # ==========================================================
    # ENVIAR OFERTA
    # ==========================================================

    if turno_para_ofrecer and turno_para_ofrecer.usuario:

        email_cliente = turno_para_ofrecer.usuario.email

        if email_cliente:

            oferta = OfertaTurno.objects.create(
                turno_cliente=turno_para_ofrecer,
                fecha_disponible=fecha_liberada
            )

            enlace = request.build_absolute_uri(
                reverse(
                    "aceptar_oferta_turno",
                    args=[oferta.token]
                )
            )

            send_mail(
                subject="¡Se liberó un turno anterior!",
                message=(
                    f"Hola {turno_para_ofrecer.usuario.get_full_name() or turno_para_ofrecer.usuario.username},\n\n"

                    "Se liberó un turno anterior al que tenés reservado.\n\n"

                    f"Nuevo horario disponible:\n"
                    f"{fecha_liberada.strftime('%d/%m/%Y %H:%M')}\n\n"

                    "Si querés adelantar tu turno, ingresá al siguiente enlace:\n\n"

                    f"QUIERO ADELANTAR MI TURNO:\n"
                    f"{enlace}\n\n"

                    "Si no querés adelantarlo, simplemente no hagas nada.\n\n"

                    "Saludos.\n"
                    "Taller Molina Henze"
                ),
                from_email=None,
                recipient_list=[email_cliente],
                fail_silently=False
            )

    messages.success(
        request,
        "Tu turno fue cancelado correctamente."
    )

    return redirect("mis_turnos")

# ==========================================================
# ACEPTAR OFERTA DE ADELANTO
# ==========================================================

def aceptar_oferta_turno(request, token):

    oferta = get_object_or_404(
        OfertaTurno,
        token=token
    )

    # ==========================================================
    # COMPROBAR SI LA OFERTA SIGUE DISPONIBLE
    # ==========================================================

    if oferta.estado != "pendiente":

        return render(
            request,
            "tallermami/oferta_no_disponible.html"
        )

    turno = oferta.turno_cliente

    # El turno fue cancelado
    if turno.estado == "cancelado":

        return render(
            request,
            "tallermami/oferta_no_disponible.html"
        )

    # ==========================================================
    # GUARDAR HORARIO ANTERIOR
    # ==========================================================

    horario_anterior = turno.fecha_turno

    # ==========================================================
    # CAMBIAR HORARIO DEL TURNO
    # ==========================================================

    turno.fecha_turno = oferta.fecha_disponible
    turno.save(
        update_fields=["fecha_turno"]
    )

    # ==========================================================
    # MARCAR OFERTA COMO ACEPTADA
    # ==========================================================

    oferta.estado = "aceptada"
    oferta.save(
        update_fields=["estado"]
    )

    # ==========================================================
    # AVISAR AL CLIENTE
    # ==========================================================

    if turno.usuario and turno.usuario.email:

        send_mail(
            subject="Turno adelantado correctamente",

            message=(
                f"Hola {turno.usuario.get_full_name() or turno.usuario.username},\n\n"

                "Tu turno fue adelantado correctamente.\n\n"

                f"Nuevo horario:\n"
                f"{turno.fecha_turno.strftime('%d/%m/%Y %H:%M')}\n\n"

                "Gracias por utilizar nuestro sistema de turnos.\n\n"

                "Taller Molina Henze"
            ),

            from_email=None,

            recipient_list=[
                turno.usuario.email
            ],

            fail_silently=True
        )

    # ==========================================================
    # MOSTRAR OFERTA ACEPTADA
    # ==========================================================

    return render(
        request,
        "tallermami/oferta_aceptada.html",
        {
            "turno": turno,
            "horario_anterior": horario_anterior
        }
    )

# ==========================================================
# ESTADO DE REPARACIÓN
# ==========================================================

def estado_reparacion(request):

    turno = None
    patente_buscada = ""

    if request.method == "POST":

        patente_buscada = request.POST.get(
            "patente",
            ""
        ).strip().upper()


        if patente_buscada:

            turno = Turno.objects.filter(

                patente__iexact=patente_buscada

            ).select_related(
                "servicio"
            ).order_by(
                "-fecha_creacion"
            ).first()


            if not turno:

                messages.error(
                    request,
                    "No encontramos ningún vehículo con esa patente."
                )


    return render(
        request,
        "tallermami/estado_reparacion.html",
        {
            "turno": turno,
            "patente_buscada": patente_buscada
        }
    )


# ==========================================================
# PDF
# ==========================================================

def descargar_pdf(request, turno_id):

    from .utils import generar_pdf_turno

    turno = get_object_or_404(
        Turno,
        id=turno_id
    )

    return generar_pdf_turno(
        turno
    )


# ==========================================================
# EXCEL
# ==========================================================

def exportar_excel(request):

    turnos = Turno.objects.select_related(
        "usuario",
        "servicio"
    ).all()

    return render(
        request,
        "tallermami/exportar_excel.html",
        {
            "turnos": turnos
        }
    )


# ==========================================================
# DESCARGAR EXCEL
# ==========================================================

def descargar_excel_turnos(request):

    from .utils import crear_excel_turnos

    archivo = crear_excel_turnos()

    response = HttpResponse(

        archivo.getvalue(),

        content_type=(
            "application/vnd.openxmlformats-officedocument."
            "spreadsheetml.sheet"
        )
    )

    response[
        "Content-Disposition"
    ] = 'attachment; filename="turnos.xlsx"'


    return response


# ==========================================================
# GUARDAR TURNO EN EXCEL
# ==========================================================

def guardar_turno_en_excel(turno):

    import os

    carpeta_excel = os.path.join(
        os.path.dirname(
            os.path.dirname(
                os.path.abspath(__file__)
            )
        ),
        "excel"
    )

    os.makedirs(
        carpeta_excel,
        exist_ok=True
    )


    ruta_excel = os.path.join(
        carpeta_excel,
        "turnos.xlsx"
    )


    # ======================================================
    # CREAR ARCHIVO SI NO EXISTE
    # ======================================================

    if os.path.exists(ruta_excel):

        libro = load_workbook(
            ruta_excel
        )

    else:

        libro = Workbook()

        hoja = libro.active

        hoja.title = "Todos los turnos"


        libro.create_sheet(
            "Clientes nuevos"
        )

        libro.create_sheet(
            "Clientes habituales"
        )


    # ======================================================
    # ASEGURAR LAS HOJAS
    # ======================================================

    if "Todos los turnos" not in libro.sheetnames:

        libro.create_sheet(
            "Todos los turnos"
        )


    if "Clientes nuevos" not in libro.sheetnames:

        libro.create_sheet(
            "Clientes nuevos"
        )


    if "Clientes habituales" not in libro.sheetnames:

        libro.create_sheet(
            "Clientes habituales"
        )


    # ======================================================
    # ENCABEZADOS
    # ======================================================

    encabezados = [

        "ID",

        "Cliente",

        "Email",

        "Teléfono",

        "Vehículo",

        "Patente",

        "Servicio",

        "Fecha",

        "Hora",

        "Tipo de cliente",

        "Estado",

        "Avance",

        "Observaciones",
    ]


    for nombre_hoja in libro.sheetnames:

        hoja = libro[nombre_hoja]

        if hoja.max_row == 1 and hoja.cell(
            1,
            1
        ).value is None:

            for columna, encabezado in enumerate(
                encabezados,
                start=1
            ):

                celda = hoja.cell(
                    row=1,
                    column=columna
                )

                celda.value = encabezado

                celda.font = Font(
                    bold=True
                )

                celda.alignment = Alignment(
                    horizontal="center"
                )


    # ======================================================
    # DATOS DEL CLIENTE
    # ======================================================

    nombre = ""

    email = ""

    telefono = ""


    if turno.usuario:

        nombre = turno.usuario.get_full_name()

        if not nombre:

            nombre = turno.usuario.username

        email = turno.usuario.email


    # ======================================================
    # TIPO DE CLIENTE
    # ======================================================

    cantidad_anterior = Turno.objects.filter(

        usuario=turno.usuario

    ).exclude(

        id=turno.id

    ).count()


    if cantidad_anterior > 0:

        tipo_cliente = "Cliente habitual"

    else:

        tipo_cliente = "Cliente nuevo"


    # ======================================================
    # DATOS
    # ======================================================

    fila = [

        turno.id,

        nombre,

        email,

        telefono,

        turno.vehiculo,

        turno.patente,

        str(turno.servicio),

        turno.fecha_turno.strftime(
            "%d/%m/%Y"
        ),

        turno.fecha_turno.strftime(
            "%H:%M"
        ),

        tipo_cliente,

        turno.get_estado_display(),

        f"{turno.porcentaje_reparacion}%",

        turno.observaciones,
    ]


    # ======================================================
    # EVITAR DUPLICADOS
    # ======================================================

    hoja_todos = libro[
        "Todos los turnos"
    ]


    ids_existentes = [

        hoja_todos.cell(
            row=fila_excel,
            column=1
        ).value

        for fila_excel in range(
            2,
            hoja_todos.max_row + 1
        )
    ]


    if turno.id not in ids_existentes:

        hoja_todos.append(
            fila
        )


    # ======================================================
    # ACTUALIZAR HOJA DE CLIENTE
    # ======================================================

    hoja_cliente = libro[
        "Clientes habituales"
        if tipo_cliente == "Cliente habitual"
        else "Clientes nuevos"
    ]


    ids_existentes_cliente = [

        hoja_cliente.cell(
            row=fila_excel,
            column=1
        ).value

        for fila_excel in range(
            2,
            hoja_cliente.max_row + 1
        )
    ]


    if turno.id not in ids_existentes_cliente:

        hoja_cliente.append(
            fila
        )


    # ======================================================
    # AJUSTAR COLUMNAS
    # ======================================================

    for nombre_hoja in libro.sheetnames:

        hoja = libro[nombre_hoja]

        for columna in range(
            1,
            hoja.max_column + 1
        ):

            letra = get_column_letter(
                columna
            )

            ancho = 15

            for celda in hoja[letra]:

                if celda.value:

                    ancho = max(
                        ancho,
                        len(str(celda.value)) + 2
                    )


            hoja.column_dimensions[
                letra
            ].width = min(
                ancho,
                40
            )


    # ======================================================
    # GUARDAR
    # ======================================================

    libro.save(
        ruta_excel
    )


# ==========================================================
# ACTUALIZAR EXCEL COMPLETO
# ==========================================================

def reconstruir_excel():

    import os

    carpeta_excel = os.path.join(
        os.path.dirname(
            os.path.dirname(
                os.path.abspath(__file__)
            )
        ),
        "excel"
    )

    os.makedirs(
        carpeta_excel,
        exist_ok=True
    )


    ruta_excel = os.path.join(
        carpeta_excel,
        "turnos.xlsx"
    )


    libro = Workbook()

    hoja_todos = libro.active

    hoja_todos.title = "Todos los turnos"

    hoja_nuevos = libro.create_sheet(
        "Clientes nuevos"
    )

    hoja_habituales = libro.create_sheet(
        "Clientes habituales"
    )


    encabezados = [

        "ID",
        "Cliente",
        "Email",
        "Vehículo",
        "Patente",
        "Servicio",
        "Fecha",
        "Hora",
        "Tipo de cliente",
        "Estado",
        "Avance",
        "Observaciones",
    ]


    for hoja in [
        hoja_todos,
        hoja_nuevos,
        hoja_habituales
    ]:

        for columna, encabezado in enumerate(
            encabezados,
            start=1
        ):

            celda = hoja.cell(
                row=1,
                column=columna
            )

            celda.value = encabezado

            celda.font = Font(
                bold=True
            )


    turnos = Turno.objects.select_related(
        "usuario",
        "servicio"
    ).all()


    for turno in turnos:

        if turno.usuario:

            cliente = (
                turno.usuario.get_full_name()
                or turno.usuario.username
            )

            email = turno.usuario.email

        else:

            cliente = "Cliente"

            email = ""


        anteriores = Turno.objects.filter(

            usuario=turno.usuario

        ).exclude(

            id=turno.id

        ).count()


        if anteriores > 0:

            tipo = "Cliente habitual"

        else:

            tipo = "Cliente nuevo"


        fila = [

            turno.id,

            cliente,

            email,

            turno.vehiculo,

            turno.patente,

            str(turno.servicio),

            turno.fecha_turno.strftime(
                "%d/%m/%Y"
            ),

            turno.fecha_turno.strftime(
                "%H:%M"
            ),

            tipo,

            turno.get_estado_display(),

            f"{turno.porcentaje_reparacion}%",

            turno.observaciones,
        ]


        hoja_todos.append(
            fila
        )


        if tipo == "Cliente nuevo":

            hoja_nuevos.append(
                fila
            )

        else:

            hoja_habituales.append(
                fila
            )


    for hoja in [
        hoja_todos,
        hoja_nuevos,
        hoja_habituales
    ]:

        for columna in range(
            1,
            hoja.max_column + 1
        ):

            letra = get_column_letter(
                columna
            )

            hoja.column_dimensions[
                letra
            ].width = 20


    libro.save(
        ruta_excel
    )