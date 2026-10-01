from io import BytesIO

from django.http import HttpResponse

from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import cm
from reportlab.platypus import (
    SimpleDocTemplate,
    Paragraph,
    Spacer,
    Table,
    TableStyle,
)

from openpyxl import Workbook
from openpyxl.styles import Font, Alignment

from .models import Turno


# ==========================================================
# GENERAR PDF
# ==========================================================

def generar_pdf_turno(turno):

    buffer = BytesIO()


    documento = SimpleDocTemplate(

        buffer,

        pagesize=A4,

        rightMargin=2 * cm,

        leftMargin=2 * cm,

        topMargin=2 * cm,

        bottomMargin=2 * cm,
    )


    estilos = getSampleStyleSheet()


    titulo = ParagraphStyle(

        "Titulo",

        parent=estilos["Title"],

        alignment=TA_CENTER,

        fontSize=20,

        spaceAfter=20
    )


    subtitulo = ParagraphStyle(

        "Subtitulo",

        parent=estilos["Heading2"],

        fontSize=13,

        spaceAfter=10
    )


    texto = ParagraphStyle(

        "Texto",

        parent=estilos["BodyText"],

        fontSize=10,

        leading=15
    )


    contenido = []


    # ======================================================
    # TÍTULO
    # ======================================================

    contenido.append(
        Paragraph(
            "TALLER MOLINA HENZE",
            titulo
        )
    )


    contenido.append(
        Paragraph(
            "Comprobante de solicitud de turno",
            subtitulo
        )
    )


    contenido.append(
        Spacer(
            1,
            10
        )
    )


    # ======================================================
    # DATOS
    # ======================================================

    usuario = turno.usuario

    if usuario:

        cliente = (
            usuario.get_full_name()
            or usuario.username
        )

        email = usuario.email

    else:

        cliente = "Cliente"

        email = ""


    datos = [

        ["Número de turno", str(turno.id)],

        ["Cliente", cliente],

        ["Email", email],

        ["Vehículo", turno.vehiculo],

        ["Patente", turno.patente],

        ["Servicio", str(turno.servicio)],

        [
            "Fecha",
            turno.fecha_turno.strftime(
                "%d/%m/%Y"
            )
        ],

        [
            "Hora",
            turno.fecha_turno.strftime(
                "%H:%M"
            )
        ],

        [
            "Estado",
            turno.get_estado_display()
        ],

        [
            "Avance",
            f"{turno.porcentaje_reparacion}%"
        ],
    ]


    tabla = Table(

        datos,

        colWidths=[
            5 * cm,
            11 * cm
        ]
    )


    tabla.setStyle(

        TableStyle([

            (
                "BACKGROUND",
                (0, 0),
                (0, -1),
                colors.HexColor("#0d2b4d")
            ),

            (
                "TEXTCOLOR",
                (0, 0),
                (0, -1),
                colors.white
            ),

            (
                "FONTNAME",
                (0, 0),
                (-1, -1),
                "Helvetica"
            ),

            (
                "FONTNAME",
                (0, 0),
                (0, -1),
                "Helvetica-Bold"
            ),

            (
                "GRID",
                (0, 0),
                (-1, -1),
                0.5,
                colors.grey
            ),

            (
                "VALIGN",
                (0, 0),
                (-1, -1),
                "MIDDLE"
            ),

            (
                "PADDING",
                (0, 0),
                (-1, -1),
                8
            ),
        ])
    )


    contenido.append(
        tabla
    )


    contenido.append(
        Spacer(
            1,
            20
        )
    )


    # ======================================================
    # OBSERVACIONES
    # ======================================================

    contenido.append(
        Paragraph(
            "Observaciones",
            subtitulo
        )
    )


    observaciones = turno.observaciones or (
        "Sin observaciones."
    )


    contenido.append(
        Paragraph(
            observaciones,
            texto
        )
    )


    contenido.append(
        Spacer(
            1,
            30
        )
    )


    contenido.append(
        Paragraph(
            "Taller Molina Henze - Servicio Integral",
            texto
        )
    )


    documento.build(
        contenido
    )


    buffer.seek(0)


    response = HttpResponse(

        buffer,

        content_type="application/pdf"
    )


    response[
        "Content-Disposition"
    ] = (
        f'attachment; '
        f'filename="turno_{turno.id}.pdf"'
    )


    return response


# ==========================================================
# CREAR EXCEL
# ==========================================================

def crear_excel_turnos():

    buffer = BytesIO()


    libro = Workbook()


    hoja_todos = libro.active

    hoja_todos.title = "Todos los turnos"


    hoja_nuevos = libro.create_sheet(
        "Clientes nuevos"
    )


    hoja_habituales = libro.create_sheet(
        "Clientes habituales"
    )


    encabezados = [

        "ID",

        "Cliente",

        "Email",

        "Vehículo",

        "Patente",

        "Servicio",

        "Fecha",

        "Hora",

        "Tipo de cliente",

        "Estado",

        "Avance",

        "Observaciones",
    ]


    for hoja in [

        hoja_todos,

        hoja_nuevos,

        hoja_habituales

    ]:

        for columna, encabezado in enumerate(
            encabezados,
            start=1
        ):

            celda = hoja.cell(

                row=1,

                column=columna
            )

            celda.value = encabezado

            celda.font = Font(
                bold=True
            )

            celda.alignment = Alignment(
                horizontal="center"
            )


    # ======================================================
    # TURNOS
    # ======================================================

    turnos = Turno.objects.select_related(
        "usuario",
        "servicio"
    ).all()


    for turno in turnos:

        if turno.usuario:

            cliente = (
                turno.usuario.get_full_name()
                or turno.usuario.username
            )

            email = turno.usuario.email

        else:

            cliente = "Cliente"

            email = ""


        cantidad_anterior = Turno.objects.filter(

            usuario=turno.usuario

        ).exclude(

            id=turno.id

        ).count()


        if cantidad_anterior > 0:

            tipo_cliente = "Cliente habitual"

        else:

            tipo_cliente = "Cliente nuevo"


        fila = [

            turno.id,

            cliente,

            email,

            turno.vehiculo,

            turno.patente,

            str(turno.servicio),

            turno.fecha_turno.strftime(
                "%d/%m/%Y"
            ),

            turno.fecha_turno.strftime(
                "%H:%M"
            ),

            tipo_cliente,

            turno.get_estado_display(),

            f"{turno.porcentaje_reparacion}%",

            turno.observaciones,
        ]


        hoja_todos.append(
            fila
        )


        if tipo_cliente == "Cliente nuevo":

            hoja_nuevos.append(
                fila
            )

        else:

            hoja_habituales.append(
                fila
            )


    # ======================================================
    # ANCHO DE COLUMNAS
    # ======================================================

    for hoja in [

        hoja_todos,

        hoja_nuevos,

        hoja_habituales

    ]:

        for columna in hoja.columns:

            maximo = 0

            letra = columna[0].column_letter


            for celda in columna:

                if celda.value:

                    longitud = len(
                        str(celda.value)
                    )

                    if longitud > maximo:

                        maximo = longitud


            hoja.column_dimensions[
                letra
            ].width = min(
                maximo + 2,
                40
            )


    libro.save(
        buffer
    )


    buffer.seek(0)


    return buffer