from django.contrib import messages
from django.contrib.auth import authenticate, login, logout
from django.contrib.auth.decorators import login_required
from django.contrib.auth.forms import AuthenticationForm
from django.contrib.auth.models import User
from django.core.mail import send_mail
from django.http import HttpResponse
from django.shortcuts import get_object_or_404, redirect, render
from django.utils import timezone
from django.contrib.auth.decorators import login_required, user_passes_test
from .models import Turno, Trabajo, OfertaTurno, Cliente, Trabajador
from .models import (
    Turno,
    Trabajo,
    OfertaTurno,
    Cliente,
    Trabajador,
    Presupuesto,
)

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

from .forms import RegistroForm, SolicitarTurnoForm
from django.urls import reverse
from .models import Turno, Trabajo, OfertaTurno, Cliente


# ==========================================================
# INICIO
# ==========================================================

def inicio(request):

    trabajos = Trabajo.objects.all()[:6]

    return render(
        request,
        "tallermami/inicio.html",
        {
            "trabajos": trabajos,
        }
    )


# ==========================================================
# LOGIN
# ==========================================================

def login_view(request):

    if request.user.is_authenticated:
        return redirect("inicio")

    if request.method == "POST":

        form = AuthenticationForm(
            request,
            data=request.POST
        )

        if form.is_valid():

            usuario = form.get_user()

            login(
                request,
                usuario
            )

            messages.success(
                request,
                "Iniciaste sesión correctamente."
            )

            return redirect("inicio")

    else:

        form = AuthenticationForm()

    return render(
        request,
        "tallermami/login.html",
        {
            "form": form
        }
    )


# ==========================================================
# REGISTRO
# ==========================================================

def registro(request):

    if request.user.is_authenticated:
        return redirect("inicio")

    if request.method == "POST":

        form = RegistroForm(request.POST)

        if form.is_valid():

            usuario = form.save()

            login(
                request,
                usuario
            )

            messages.success(
                request,
                "Tu cuenta fue creada correctamente."
            )

            return redirect("inicio")

    else:

        form = RegistroForm()

    return render(
        request,
        "tallermami/registro.html",
        {
            "form": form
        }
    )


# ==========================================================
# LOGOUT
# ==========================================================

def logout_view(request):

    logout(request)

    messages.success(
        request,
        "Sesión cerrada correctamente."
    )

    return redirect("inicio")


# ==========================================================
# SOLICITAR TURNO
# ==========================================================


def solicitar_turno(request):
    if request.method == "POST":
        form = SolicitarTurnoForm(request.POST)

        if form.is_valid():
            # Obtener los datos del cliente
            nombre = form.cleaned_data["nombre"].strip()
            apellido = form.cleaned_data["apellido"].strip()
            email_cliente = form.cleaned_data["email"].strip()
            telefono = form.cleaned_data["telefono"].strip()

            # ==============================================
            # CREAR O ACTUALIZAR EL CLIENTE AUTOMÁTICAMENTE
            # ==============================================

            cliente = Cliente.objects.filter(
                email__iexact=email_cliente
            ).first()

            if cliente:
                cliente.nombre = nombre
                cliente.apellido = apellido
                cliente.email = email_cliente
                cliente.telefono = telefono
                cliente.save()
            else:
                Cliente.objects.create(
                    nombre=nombre,
                    apellido=apellido,
                    email=email_cliente,
                    telefono=telefono
                )

            # ==============================================
            # ASOCIAR EL TURNO A UN USUARIO
            # ==============================================

            if request.user.is_authenticated:
                usuario = request.user

                # Usar el correo ingresado en la solicitud
                # para que las notificaciones lleguen al cliente.
                if usuario.email != email_cliente:
                    usuario.email = email_cliente
                    usuario.save(update_fields=["email"])

            else:
                # Buscar una cuenta existente por correo
                usuario = User.objects.filter(
                    email__iexact=email_cliente
                ).first()

                if not usuario:
                    # Crear un nombre de usuario único
                    base_username = email_cliente.split("@")[0]
                    username = base_username
                    contador = 1

                    while User.objects.filter(
                        username=username
                    ).exists():
                        username = f"{base_username}{contador}"
                        contador += 1

                    usuario = User(
                        username=username,
                        email=email_cliente
                    )
                    usuario.set_unusable_password()
                    usuario.save()

            # ==============================================
            # GUARDAR EL TURNO
            # ==============================================

            turno = form.save(commit=False)
            turno.usuario = usuario
            turno.save()

            # ==============================================
            # GUARDAR EN EXCEL
            # ==============================================

            try:
                guardar_turno_en_excel(turno)
            except Exception as error:
                print(
                    f"Error al guardar el turno en Excel: {error}"
                )

            # ==============================================
            # ENVIAR EMAIL DE CONFIRMACIÓN
            # ==============================================

            try:
                send_mail(
                    subject="Solicitud de turno - Taller Molina Henze",
                    message=(
                        f"Hola {nombre},\n\n"
                        "Recibimos correctamente tu solicitud de turno.\n\n"
                        f"Vehículo: {turno.vehiculo}\n"
                        f"Patente: {turno.patente}\n"
                        f"Servicio: {turno.servicio}\n"
                        f"Fecha: {turno.fecha_turno.strftime('%d/%m/%Y %H:%M')}\n\n"
                        "Te esperamos en Taller Molina Henze.\n\n"
                        "Saludos."
                    ),
                    from_email=None,
                    recipient_list=[email_cliente],
                    fail_silently=False
                )
            except Exception as error:
                print(f"Error al enviar email: {error}")

            return redirect(
                "turno_exitoso",
                turno_id=turno.id
            )

    else:
        form = SolicitarTurnoForm()

    return render(
        request,
        "tallermami/solicitar_turno.html",
        {"form": form}
    )

# ==========================================================
# TURNO EXITOSO
# ==========================================================

def turno_exitoso(request, turno_id):

    turno = get_object_or_404(
        Turno,
        id=turno_id
    )

    return render(
        request,
        "tallermami/turno_exitoso.html",
        {
            "turno": turno
        }
    )


# ==========================================================
# MIS TURNOS
# ==========================================================




@login_required
def mis_turnos(request):

    turnos = Turno.objects.filter(
        usuario=request.user
    ).exclude(
        estado="cancelado"
    ).select_related(
        "servicio"
    )

    return render(
        request,
        "tallermami/mis_turnos.html",
        {
            "turnos": turnos
        }
    )



# ==========================================================
# CANCELAR TURNO
# ==========================================================


@login_required
def cancelar_turno(request, turno_id):

    turno = get_object_or_404(
        Turno,
        id=turno_id,
        usuario=request.user
    )

    # Evitar cancelar un turno que ya fue cancelado
    if turno.estado == "cancelado":
        messages.error(
            request,
            "Este turno ya está cancelado."
        )
        return redirect("mis_turnos")

    # Guardar el horario que quedará libre
    fecha_liberada = turno.fecha_turno

    # Cancelar el turno
    turno.estado = "cancelado"
    turno.save(update_fields=["estado"])

    # Solo ofrecer horarios que todavía no pasaron
    if fecha_liberada <= timezone.now():
        messages.success(
            request,
            "Tu turno fue cancelado correctamente."
        )
        return redirect("mis_turnos")

    # Buscar clientes con turnos posteriores al horario liberado
    candidatos = (
        Turno.objects
        .filter(
            fecha_turno__gt=fecha_liberada,
            estado="pendiente",
            usuario__isnull=False
        )
        .exclude(id=turno.id)
        .exclude(usuario__email="")
        .select_related("usuario", "servicio")
        .order_by("fecha_turno")
    )

    # Elegir el primer cliente elegible que no tenga
    # otra oferta pendiente para su turno
    turno_para_ofrecer = None

    for candidato in candidatos:
        tiene_oferta_pendiente = OfertaTurno.objects.filter(
            turno_cliente=candidato,
            estado="pendiente"
        ).exists()

        if not tiene_oferta_pendiente:
            turno_para_ofrecer = candidato
            break

    # Crear y enviar la oferta
    if turno_para_ofrecer:

        oferta = OfertaTurno.objects.create(
            turno_cliente=turno_para_ofrecer,
            fecha_disponible=fecha_liberada
        )

        enlace = request.build_absolute_uri(
            reverse(
                "aceptar_oferta_turno",
                args=[oferta.token]
            )
        )

        send_mail(
            subject="¡Podés adelantar tu turno!",
            message=(
                f"Hola {turno_para_ofrecer.usuario.get_full_name() or turno_para_ofrecer.usuario.username},\n\n"
                "Se liberó un horario anterior al de tu turno actual.\n\n"
                "Si querés adelantar tu turno, ingresá al siguiente enlace:\n\n"
                "QUIERO ADELANTAR MI TURNO:\n"
                f"{enlace}\n\n"
                "La disponibilidad está sujeta a confirmación.\n\n"
                "Saludos.\n"
                "Taller Molina Henze"
            ),
            from_email=None,
            recipient_list=[
                turno_para_ofrecer.usuario.email
            ],
            fail_silently=False
        )

    messages.success(
        request,
        "Tu turno fue cancelado correctamente."
    )

    return redirect("mis_turnos")

# ==========================================================
# ACEPTAR OFERTA DE ADELANTO
# ==========================================================

def aceptar_oferta_turno(request, token):

    oferta = get_object_or_404(
        OfertaTurno,
        token=token
    )

    # ==========================================================
    # COMPROBAR SI LA OFERTA SIGUE DISPONIBLE
    # ==========================================================

    if oferta.estado != "pendiente":

        return render(
            request,
            "tallermami/oferta_no_disponible.html"
        )

    turno = oferta.turno_cliente

    # El turno fue cancelado
    if turno.estado == "cancelado":

        return render(
            request,
            "tallermami/oferta_no_disponible.html"
        )

    # ==========================================================
    # GUARDAR HORARIO ANTERIOR
    # ==========================================================

    horario_anterior = turno.fecha_turno

    # ==========================================================
    # CAMBIAR HORARIO DEL TURNO
    # ==========================================================

    turno.fecha_turno = oferta.fecha_disponible
    turno.save(
        update_fields=["fecha_turno"]
    )

    # ==========================================================
    # MARCAR OFERTA COMO ACEPTADA
    # ==========================================================

    oferta.estado = "aceptada"
    oferta.save(
        update_fields=["estado"]
    )

    # ==========================================================
    # AVISAR AL CLIENTE
    # ==========================================================

    if turno.usuario and turno.usuario.email:

        send_mail(
            subject="Turno adelantado correctamente",

            message=(
                f"Hola {turno.usuario.get_full_name() or turno.usuario.username},\n\n"

                "Tu turno fue adelantado correctamente.\n\n"

                f"Nuevo horario:\n"
                f"{turno.fecha_turno.strftime('%d/%m/%Y %H:%M')}\n\n"

                "Gracias por utilizar nuestro sistema de turnos.\n\n"

                "Taller Molina Henze"
            ),

            from_email=None,

            recipient_list=[
                turno.usuario.email
            ],

            fail_silently=True
        )

    # ==========================================================
    # MOSTRAR OFERTA ACEPTADA
    # ==========================================================

    return render(
        request,
        "tallermami/oferta_aceptada.html",
        {
            "turno": turno,
            "horario_anterior": horario_anterior
        }
    )

# ==========================================================
# ESTADO DE REPARACIÓN
# ==========================================================

def estado_reparacion(request):

    turno = None
    patente_buscada = ""

    if request.method == "POST":

        patente_buscada = request.POST.get(
            "patente",
            ""
        ).strip().upper()


        if patente_buscada:

            turno = Turno.objects.filter(

                patente__iexact=patente_buscada

            ).select_related(
                "servicio"
            ).order_by(
                "-fecha_creacion"
            ).first()


            if not turno:

                messages.error(
                    request,
                    "No encontramos ningún vehículo con esa patente."
                )


    return render(
        request,
        "tallermami/estado_reparacion.html",
        {
            "turno": turno,
            "patente_buscada": patente_buscada
        }
    )


# ==========================================================
# PDF
# ==========================================================

def descargar_pdf(request, turno_id):

    from .utils import generar_pdf_turno

    turno = get_object_or_404(
        Turno,
        id=turno_id
    )

    return generar_pdf_turno(
        turno
    )


# ==========================================================
# EXCEL
# ==========================================================

def exportar_excel(request):

    turnos = Turno.objects.select_related(
        "usuario",
        "servicio"
    ).all()

    return render(
        request,
        "tallermami/exportar_excel.html",
        {
            "turnos": turnos
        }
    )


# ==========================================================
# DESCARGAR EXCEL
# ==========================================================


@login_required
def descargar_excel_turnos(request):
    # Permitir la descarga únicamente a los cargos autorizados.
    if not (es_jefe_taller(request.user) or es_secretario(request.user)):
        return HttpResponse("No tenés permiso para descargar este archivo.", status=403)

    from .utils import crear_excel_turnos

    turnos, busqueda = obtener_turnos_filtrados(request)
    archivo = crear_excel_turnos(turnos)

    response = HttpResponse(
        archivo.getvalue(),
        content_type=(
            "application/vnd.openxmlformats-officedocument."
            "spreadsheetml.sheet"
        ),
    )

    response["Content-Disposition"] = 'attachment; filename="turnos.xlsx"'
    return response


# ==========================================================
# GUARDAR TURNO EN EXCEL
# ==========================================================

def guardar_turno_en_excel(turno):

    import os

    carpeta_excel = os.path.join(
        os.path.dirname(
            os.path.dirname(
                os.path.abspath(__file__)
            )
        ),
        "excel"
    )

    os.makedirs(
        carpeta_excel,
        exist_ok=True
    )


    ruta_excel = os.path.join(
        carpeta_excel,
        "turnos.xlsx"
    )


    # ======================================================
    # CREAR ARCHIVO SI NO EXISTE
    # ======================================================

    if os.path.exists(ruta_excel):

        libro = load_workbook(
            ruta_excel
        )

    else:

        libro = Workbook()

        hoja = libro.active

        hoja.title = "Todos los turnos"


        libro.create_sheet(
            "Clientes nuevos"
        )

        libro.create_sheet(
            "Clientes habituales"
        )


    # ======================================================
    # ASEGURAR LAS HOJAS
    # ======================================================

    if "Todos los turnos" not in libro.sheetnames:

        libro.create_sheet(
            "Todos los turnos"
        )


    if "Clientes nuevos" not in libro.sheetnames:

        libro.create_sheet(
            "Clientes nuevos"
        )


    if "Clientes habituales" not in libro.sheetnames:

        libro.create_sheet(
            "Clientes habituales"
        )


    # ======================================================
    # ENCABEZADOS
    # ======================================================

    encabezados = [

        "ID",

        "Cliente",

        "Email",

        "Teléfono",

        "Vehículo",

        "Patente",

        "Servicio",

        "Fecha",

        "Hora",

        "Tipo de cliente",

        "Estado",

        "Avance",

        "Observaciones",
    ]


    for nombre_hoja in libro.sheetnames:

        hoja = libro[nombre_hoja]

        if hoja.max_row == 1 and hoja.cell(
            1,
            1
        ).value is None:

            for columna, encabezado in enumerate(
                encabezados,
                start=1
            ):

                celda = hoja.cell(
                    row=1,
                    column=columna
                )

                celda.value = encabezado

                celda.font = Font(
                    bold=True
                )

                celda.alignment = Alignment(
                    horizontal="center"
                )


    # ======================================================
    # DATOS DEL CLIENTE
    # ======================================================

    nombre = ""

    email = ""

    telefono = ""


    if turno.usuario:

        nombre = turno.usuario.get_full_name()

        if not nombre:

            nombre = turno.usuario.username

        email = turno.usuario.email


    # ======================================================
    # TIPO DE CLIENTE
    # ======================================================

    cantidad_anterior = Turno.objects.filter(

        usuario=turno.usuario

    ).exclude(

        id=turno.id

    ).count()


    if cantidad_anterior > 0:

        tipo_cliente = "Cliente habitual"

    else:

        tipo_cliente = "Cliente nuevo"


    # ======================================================
    # DATOS
    # ======================================================

    fila = [

        turno.id,

        nombre,

        email,

        telefono,

        turno.vehiculo,

        turno.patente,

        str(turno.servicio),

        turno.fecha_turno.strftime(
            "%d/%m/%Y"
        ),

        turno.fecha_turno.strftime(
            "%H:%M"
        ),

        tipo_cliente,

        turno.get_estado_display(),

        f"{turno.porcentaje_reparacion}%",

        turno.observaciones,
    ]


    # ======================================================
    # EVITAR DUPLICADOS
    # ======================================================

    hoja_todos = libro[
        "Todos los turnos"
    ]


    ids_existentes = [

        hoja_todos.cell(
            row=fila_excel,
            column=1
        ).value

        for fila_excel in range(
            2,
            hoja_todos.max_row + 1
        )
    ]


    if turno.id not in ids_existentes:

        hoja_todos.append(
            fila
        )


    # ======================================================
    # ACTUALIZAR HOJA DE CLIENTE
    # ======================================================

    hoja_cliente = libro[
        "Clientes habituales"
        if tipo_cliente == "Cliente habitual"
        else "Clientes nuevos"
    ]


    ids_existentes_cliente = [

        hoja_cliente.cell(
            row=fila_excel,
            column=1
        ).value

        for fila_excel in range(
            2,
            hoja_cliente.max_row + 1
        )
    ]


    if turno.id not in ids_existentes_cliente:

        hoja_cliente.append(
            fila
        )


    # ======================================================
    # AJUSTAR COLUMNAS
    # ======================================================

    for nombre_hoja in libro.sheetnames:

        hoja = libro[nombre_hoja]

        for columna in range(
            1,
            hoja.max_column + 1
        ):

            letra = get_column_letter(
                columna
            )

            ancho = 15

            for celda in hoja[letra]:

                if celda.value:

                    ancho = max(
                        ancho,
                        len(str(celda.value)) + 2
                    )


            hoja.column_dimensions[
                letra
            ].width = min(
                ancho,
                40
            )


    # ======================================================
    # GUARDAR
    # ======================================================

    libro.save(
        ruta_excel
    )


# ==========================================================
# ACTUALIZAR EXCEL COMPLETO
# ==========================================================

def reconstruir_excel():

    import os

    carpeta_excel = os.path.join(
        os.path.dirname(
            os.path.dirname(
                os.path.abspath(__file__)
            )
        ),
        "excel"
    )

    os.makedirs(
        carpeta_excel,
        exist_ok=True
    )


    ruta_excel = os.path.join(
        carpeta_excel,
        "turnos.xlsx"
    )


    libro = Workbook()

    hoja_todos = libro.active

    hoja_todos.title = "Todos los turnos"

    hoja_nuevos = libro.create_sheet(
        "Clientes nuevos"
    )

    hoja_habituales = libro.create_sheet(
        "Clientes habituales"
    )


    encabezados = [

        "ID",
        "Cliente",
        "Email",
        "Vehículo",
        "Patente",
        "Servicio",
        "Fecha",
        "Hora",
        "Tipo de cliente",
        "Estado",
        "Avance",
        "Observaciones",
    ]


    for hoja in [
        hoja_todos,
        hoja_nuevos,
        hoja_habituales
    ]:

        for columna, encabezado in enumerate(
            encabezados,
            start=1
        ):

            celda = hoja.cell(
                row=1,
                column=columna
            )

            celda.value = encabezado

            celda.font = Font(
                bold=True
            )


    turnos = Turno.objects.select_related(
        "usuario",
        "servicio"
    ).all()


    for turno in turnos:

        if turno.usuario:

            cliente = (
                turno.usuario.get_full_name()
                or turno.usuario.username
            )

            email = turno.usuario.email

        else:

            cliente = "Cliente"

            email = ""


        anteriores = Turno.objects.filter(

            usuario=turno.usuario

        ).exclude(

            id=turno.id

        ).count()


        if anteriores > 0:

            tipo = "Cliente habitual"

        else:

            tipo = "Cliente nuevo"


        fila = [

            turno.id,

            cliente,

            email,

            turno.vehiculo,

            turno.patente,

            str(turno.servicio),

            turno.fecha_turno.strftime(
                "%d/%m/%Y"
            ),

            turno.fecha_turno.strftime(
                "%H:%M"
            ),

            tipo,

            turno.get_estado_display(),

            f"{turno.porcentaje_reparacion}%",

            turno.observaciones,
        ]


        hoja_todos.append(
            fila
        )


        if tipo == "Cliente nuevo":

            hoja_nuevos.append(
                fila
            )

        else:

            hoja_habituales.append(
                fila
            )


    for hoja in [
        hoja_todos,
        hoja_nuevos,
        hoja_habituales
    ]:

        for columna in range(
            1,
            hoja.max_column + 1
        ):

            letra = get_column_letter(
                columna
            )

            hoja.column_dimensions[
                letra
            ].width = 20


    libro.save(
        ruta_excel
    )
   
# ==========================================================
# PERMISOS DE LOS CARGOS
# ==========================================================


# ==========================================================
# BUSCADOR COMPARTIDO PARA JEFE Y SECRETARIO
# ==========================================================

from django.db.models import Q


def obtener_turnos_filtrados(request):
    busqueda = request.GET.get("buscar", "").strip()

    turnos = Turno.objects.select_related(
        "usuario",
        "servicio",
        "trabajador"
    ).order_by("-fecha_turno")

    if busqueda:
        # Buscar también entre los clientes registrados.
        emails_clientes = Cliente.objects.filter(
            Q(nombre__icontains=busqueda)
            | Q(apellido__icontains=busqueda)
            | Q(email__icontains=busqueda)
        ).values_list("email", flat=True)

        turnos = turnos.filter(
            Q(patente__icontains=busqueda)
            | Q(vehiculo__icontains=busqueda)
            | Q(usuario__username__icontains=busqueda)
            | Q(usuario__first_name__icontains=busqueda)
            | Q(usuario__last_name__icontains=busqueda)
            | Q(usuario__email__icontains=busqueda)
            | Q(usuario__email__in=emails_clientes)
        ).distinct()

    return turnos, busqueda


# ==========================================================
# PERMISOS DE LOS CARGOS
# ==========================================================

def es_jefe_taller(user):
    return (
        user.is_authenticated
        and (
            user.is_superuser
            or user.groups.filter(name="Jefe de taller").exists()
        )
    )


def es_secretario(user):
    return (
        user.is_authenticated
        and (
            user.is_superuser
            or user.groups.filter(name="Secretario").exists()
        )
    )

# ==========================================================
# VISTA DEL JEFE DE TALLER
# ==========================================================

@login_required
@user_passes_test(es_jefe_taller)
def vista_jefe_taller(request):
    trabajos = Trabajo.objects.all().order_by("nombre")

    trabajadores = Trabajador.objects.all().order_by(
        "apellido", "nombre"
    )

    clientes = Cliente.objects.all().order_by(
        "apellido", "nombre"
    )

    turnos, busqueda = obtener_turnos_filtrados(request)
    
    return render(
        request,
        "tallermami/jefe_taller.html",
        {
            "trabajos": trabajos,
            "trabajadores": trabajadores,
            "clientes": clientes,
            "turnos": turnos,
            "busqueda": busqueda,
        }
    )


# ==========================================================
# VISTA DEL SECRETARIO
# ==========================================================

@login_required
@user_passes_test(es_secretario)
def vista_secretario(request):
    clientes = Cliente.objects.all().order_by(
        "apellido", "nombre"
    )

    turnos, busqueda = obtener_turnos_filtrados(request)

    return render(
        request,
        "tallermami/secretario.html",
        {
            "clientes": clientes,
            "turnos": turnos,
            "busqueda": busqueda,
        }
    )

# ==========================================================
# SECCIONES DEL JEFE DE TALLER
# ==========================================================

@login_required
@user_passes_test(es_jefe_taller)
def jefe_trabajos(request):
    return render(request, "tallermami/seccion_panel.html", {
        "panel": "jefe",
        "seccion": "trabajos",
        "titulo": "Trabajos y servicios",
        "trabajos": Trabajo.objects.all().order_by("nombre"),
    })


@login_required
@user_passes_test(es_jefe_taller)
def jefe_trabajadores(request):
    return render(request, "tallermami/seccion_panel.html", {
        "panel": "jefe",
        "seccion": "trabajadores",
        "titulo": "Trabajadores",
        "trabajadores": Trabajador.objects.all().order_by(
            "apellido", "nombre"
        ),
    })


@login_required
@user_passes_test(es_jefe_taller)
def jefe_clientes(request):
    return render(request, "tallermami/seccion_panel.html", {
        "panel": "jefe",
        "seccion": "clientes",
        "titulo": "Clientes",
        "clientes": Cliente.objects.all().order_by("apellido", "nombre"),
    })


@login_required
@user_passes_test(es_jefe_taller)
def jefe_turnos(request):
    turnos, busqueda = obtener_turnos_filtrados(request)
    return render(request, "tallermami/seccion_panel.html", {
        "panel": "jefe",
        "seccion": "turnos",
        "titulo": "Turnos",
        "turnos": turnos,
        "busqueda": busqueda,
    })


# ==========================================================
# SECCIONES DEL SECRETARIO
# ==========================================================

@login_required
@user_passes_test(es_secretario)
def secretario_clientes(request):
    return render(request, "tallermami/seccion_panel.html", {
        "panel": "secretario",
        "seccion": "clientes",
        "titulo": "Clientes",
        "clientes": Cliente.objects.all().order_by("apellido", "nombre"),
    })


@login_required
@user_passes_test(es_secretario)
def secretario_turnos(request):
    turnos, busqueda = obtener_turnos_filtrados(request)
    return render(request, "tallermami/seccion_panel.html", {
        "panel": "secretario",
        "seccion": "turnos",
        "titulo": "Turnos",
        "turnos": turnos,
        "busqueda": busqueda,
    })
# ==========================================================
# RESPONDER PRESUPUESTO POR CORREO
# ==========================================================

def responder_presupuesto(request, token, decision):

    if decision not in ("aceptar", "rechazar"):
        return HttpResponse("Decisión inválida.", status=400)

    presupuesto = get_object_or_404(
        Presupuesto.objects.select_related(
            "turno",
            "turno__usuario",
        ),
        token=token,
    )

    turno = presupuesto.turno

    # Si ya respondió, no se permite cambiar la decisión.
    if presupuesto.estado != "pendiente":
        return render(
            request,
            "tallermami/presupuesto_responder.html",
            {
                "presupuesto": presupuesto,
                "decision": decision,
                "puede_responder": False,
                "resultado": (
                    "Este presupuesto ya fue respondido. "
                    "No se puede modificar la decisión."
                ),
            },
        )

    # El enlace solo muestra una confirmación.
    # La decisión se registra al enviar el formulario.
    if request.method == "POST":

        if request.POST.get("confirmar") != "si":
            return HttpResponse(
                "Debés confirmar tu decisión.",
                status=400,
            )

        if decision == "aceptar":
            presupuesto.estado = "aceptado"
            turno.estado = "presupuesto_aceptado"
            resultado = "Aceptaste el presupuesto correctamente."

        else:
            presupuesto.estado = "rechazado"
            turno.estado = "presupuesto_rechazado"
            resultado = "Rechazaste el presupuesto."

        presupuesto.fecha_respuesta = timezone.now()
        presupuesto.save(
            update_fields=["estado", "fecha_respuesta"]
        )

        turno.save(update_fields=["estado"])

        return render(
            request,
            "tallermami/presupuesto_responder.html",
            {
                "presupuesto": presupuesto,
                "decision": decision,
                "puede_responder": False,
                "resultado": resultado,
            },
        )

    return render(
        request,
        "tallermami/presupuesto_responder.html",
        {
            "presupuesto": presupuesto,
            "decision": decision,
            "puede_responder": True,
            "resultado": "",
        },
    )